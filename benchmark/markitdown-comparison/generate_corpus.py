#!/usr/bin/env python3
"""Generate the benchmark corpus from the ground-truth Markdown documents.

Ground truth lives in ground_truth/*.md (a small Markdown subset: h1-h3,
single-line paragraphs, '- ' bullets, '1. ' numbered items, pipe tables,
[text](url) links, **bold**, *italic*). Each document is rendered into the
formats listed for it in PLAN below. Output goes to corpus/ plus
corpus/manifest.json, which records for every file its ground truth and which
structural elements that format can carry (used by score.py so impossible
elements are never penalised).

Deterministic: no randomness, fixed document metadata, normalised ZIP
timestamps. The only non-deterministic bytes are PDF creation dates written by
Chrome, so the committed corpus is the canonical copy.

Usage: python generate_corpus.py [--only docx,html,...] [--skip-pdf]
"""
import argparse
import html
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
GT = HERE / "ground_truth"
OUT = HERE / "corpus"
CHROME = os.environ.get(
    "CHROME_BIN", "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
)
FIXED = datetime(2026, 10, 25, 0, 0, 0)

# document -> (language tag, formats)
PLAN = {
    "01_handbook_en": ("en", ["docx", "html", "pdf"]),
    "02_sales_report": ("en", ["docx", "html", "pdf"]),
    "03_turkce_rapor": ("tr", ["docx", "html", "pdf"]),
    "04_multilingual": ("mul", ["docx", "html", "pdf"]),
    "05_long_report": ("en", ["docx", "html", "pdf"]),
    "06_deck_roadmap": ("en", ["pptx"]),
    "07_deck_turkce": ("tr", ["pptx"]),
    "08_workbook_inventory": ("en", ["xlsx", "csv"]),
    "09_workbook_turkce": ("tr", ["xlsx", "csv"]),
    "10_workbook_multilingual": ("mul", ["xlsx", "csv"]),
}

# Elements each format can carry (see README "Method"). PDF structure is
# recovered heuristically from layout, but the information is visually
# present, so it is scored; PPTX has slide titles but no heading levels.
CARRIES = {
    "docx": ["text", "order", "headings", "heading_levels", "lists", "tables", "links"],
    "html": ["text", "order", "headings", "heading_levels", "lists", "tables", "links"],
    "pdf": ["text", "order", "headings", "heading_levels", "lists", "tables", "links"],
    "pptx": ["text", "order", "headings", "lists", "tables", "links"],
    "xlsx": ["text", "order", "tables"],
    "csv": ["text", "order", "tables"],
}

# ----------------------------------------------------------------- parsing
INLINE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|\*([^*]+)\*")


def runs(text):
    """Split inline Markdown into (text, bold, italic, url) runs."""
    out, pos = [], 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], False, False, None))
        if m.group(1):
            out.append((m.group(1), False, False, m.group(2)))
        elif m.group(3):
            out.append((m.group(3), True, False, None))
        else:
            out.append((m.group(4), False, True, None))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False, False, None))
    return out


def parse_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def parse_md(text):
    blocks, lines, i = [], text.splitlines(), 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1
        elif ln.startswith("#"):
            lvl = len(ln) - len(ln.lstrip("#"))
            blocks.append(("h", lvl, ln[lvl:].strip()))
            i += 1
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.fullmatch(r"\|[\s\-|:]+\|", lines[i]):
                    rows.append(parse_row(lines[i]))
                i += 1
            blocks.append(("table", rows))
        elif ln.startswith("- "):
            items = []
            while i < len(lines) and lines[i].startswith("- "):
                items.append(lines[i][2:].strip())
                i += 1
            blocks.append(("ul", items))
        elif re.match(r"\d+\. ", ln):
            items = []
            while i < len(lines) and re.match(r"\d+\. ", lines[i]):
                items.append(re.sub(r"^\d+\. ", "", lines[i]).strip())
                i += 1
            blocks.append(("ol", items))
        else:
            blocks.append(("p", ln.strip()))
            i += 1
    return blocks


# ----------------------------------------------------------------- helpers
def normalize_zip(path):
    """Rewrite an OOXML zip with fixed timestamps (byte-stable output)."""
    tmp = str(path) + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 10, 25, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            data = zin.read(info.filename)
            if info.filename == "docProps/core.xml":  # openpyxl stamps "modified" with the wall clock on save
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>2026-10-25T00:00:00Z", data)
            zout.writestr(zi, data)
    os.replace(tmp, path)


def to_number(s):
    if re.fullmatch(r"-?\d+", s):
        return int(s)
    if re.fullmatch(r"-?\d+\.\d+", s):
        return float(s)
    return None


# -------------------------------------------------------------------- DOCX
def write_docx(blocks, path, lang):
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    doc = Document()
    cp = doc.core_properties
    cp.author, cp.last_modified_by, cp.title = "Benchmark", "Benchmark", path.stem
    cp.created = cp.modified = FIXED

    def add_runs(par, text):
        for t, bold, ital, url in runs(text):
            if url:
                rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
                h = OxmlElement("w:hyperlink")
                h.set(qn("r:id"), rid)
                r = OxmlElement("w:r")
                rpr = OxmlElement("w:rPr")
                col = OxmlElement("w:color"); col.set(qn("w:val"), "0563C1"); rpr.append(col)
                u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
                r.append(rpr)
                te = OxmlElement("w:t"); te.text = t
                te.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                r.append(te); h.append(r)
                par._p.append(h)
            else:
                run = par.add_run(t)
                run.bold, run.italic = bold or None, ital or None

    def restart_numbering():
        numbering = doc.part.numbering_part.numbering_definitions._numbering
        num_id = doc.styles["List Number"].element.pPr.numPr.numId.val
        abstract = numbering.num_having_numId(num_id).abstractNumId.val
        new = numbering.add_num(abstract)
        new.add_lvlOverride(ilvl=0).add_startOverride(1)
        return new.numId

    for b in blocks:
        if b[0] == "h":
            h = doc.add_heading(level=b[1])
            add_runs(h, b[2])
        elif b[0] == "p":
            add_runs(doc.add_paragraph(), b[1])
        elif b[0] == "ul":
            for it in b[1]:
                add_runs(doc.add_paragraph(style="List Bullet"), it)
        elif b[0] == "ol":
            nid = restart_numbering()
            for it in b[1]:
                p = doc.add_paragraph(style="List Number")
                npr = p._p.get_or_add_pPr().get_or_add_numPr()
                npr.get_or_add_ilvl().val = 0
                npr.get_or_add_numId().val = nid
                add_runs(p, it)
        elif b[0] == "table":
            rows = b[1]
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, val in enumerate(row):
                    cell = t.cell(ri, ci)
                    if val:
                        run = cell.paragraphs[0].add_run(val)
                        run.bold = True if ri == 0 else None
            doc.add_paragraph()
    doc.save(path)
    normalize_zip(path)


# -------------------------------------------------------------------- HTML
CSS = """
body{font-family:"Helvetica Neue",Arial,"Hiragino Sans","PingFang SC",sans-serif;font-size:11pt;line-height:1.5;max-width:46em;margin:2em auto;padding:0 1em;color:#111}
h1{font-size:22pt;margin:1em 0 .4em}h2{font-size:16pt;margin:1.2em 0 .4em}h3{font-size:13pt;margin:1em 0 .3em}
table{border-collapse:collapse;margin:1em 0;width:100%}th,td{border:1px solid #888;padding:4px 8px;text-align:left}th{background:#eee}
h1,h2,h3{break-after:avoid}table,tr{break-inside:avoid}
@page{size:A4;margin:20mm}
"""


def inline_html(text):
    out = []
    for t, bold, ital, url in runs(text):
        s = html.escape(t)
        if bold:
            s = f"<strong>{s}</strong>"
        if ital:
            s = f"<em>{s}</em>"
        if url:
            s = f'<a href="{html.escape(url)}">{s}</a>'
        out.append(s)
    return "".join(out)


def render_html(blocks, lang, title):
    body = []
    for b in blocks:
        if b[0] == "h":
            body.append(f"<h{b[1]}>{inline_html(b[2])}</h{b[1]}>")
        elif b[0] == "p":
            body.append(f"<p>{inline_html(b[1])}</p>")
        elif b[0] in ("ul", "ol"):
            body.append(f"<{b[0]}>" + "".join(f"<li>{inline_html(i)}</li>" for i in b[1]) + f"</{b[0]}>")
        elif b[0] == "table":
            rows = b[1]
            head = "".join(f"<th>{inline_html(c)}</th>" for c in rows[0])
            trs = "".join("<tr>" + "".join(f"<td>{inline_html(c)}</td>" for c in r) + "</tr>" for r in rows[1:])
            body.append(f"<table><thead><tr>{head}</tr></thead><tbody>{trs}</tbody></table>")
    htmllang = {"mul": "en"}.get(lang, lang)
    return (f'<!doctype html>\n<html lang="{htmllang}"><head><meta charset="utf-8">'
            f"<title>{html.escape(title)}</title><style>{CSS}</style></head><body>\n"
            + "\n".join(body) + "\n</body></html>\n")


def write_pdf(html_path, pdf_path):
    # Remove the committed file first so a failed Chrome run can never pass off a stale PDF as new.
    Path(pdf_path).unlink(missing_ok=True)
    with tempfile.TemporaryDirectory() as prof:
        cmd = [CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
               f"--user-data-dir={prof}", f"--print-to-pdf={pdf_path}", f"file://{html_path}"]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
    if r.returncode != 0 or not Path(pdf_path).exists():
        raise RuntimeError(f"Chrome did not produce {pdf_path} (exit {r.returncode}): {r.stderr[-500:]}")


# -------------------------------------------------------------------- PPTX
def write_pptx(blocks, path):
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.core_properties.author = prs.core_properties.last_modified_by = "Benchmark"
    prs.core_properties.created = prs.core_properties.modified = FIXED
    prs.core_properties.title = path.stem

    def add_runs(par, text):
        for t, bold, ital, url in runs(text):
            r = par.add_run()
            r.text = t
            if bold:
                r.font.bold = True
            if ital:
                r.font.italic = True
            if url:
                r.hyperlink.address = url

    # group blocks into slides: h1 -> title slide, each h2 -> one slide
    slides, cur = [], None
    for b in blocks:
        if b[0] == "h" and b[1] == 1:
            slides.append({"title": b[2], "title_slide": True, "body": []})
            cur = None
        elif b[0] == "h" and b[1] == 2:
            cur = {"title": b[2], "title_slide": False, "body": []}
            slides.append(cur)
        elif cur is not None:
            cur["body"].append(b)
    for s in slides:
        bullets = [x for blk in s["body"] if blk[0] in ("ul", "ol") for x in blk[1]]
        tables = [blk[1] for blk in s["body"] if blk[0] == "table"]
        if s["title_slide"]:
            sl = prs.slides.add_slide(prs.slide_layouts[0])
            sl.shapes.title.text = s["title"]
            sub = sl.placeholders[1]
            sub._element.getparent().remove(sub._element)
            continue
        layout = prs.slide_layouts[1] if bullets else prs.slide_layouts[5]
        sl = prs.slides.add_slide(layout)
        sl.shapes.title.text = s["title"]
        if bullets:
            tf = sl.placeholders[1].text_frame
            for i, it in enumerate(bullets):
                par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                add_runs(par, it)
                for r in par.runs:
                    r.font.size = Pt(24)
        for rows in tables:
            top = Inches(1.9)
            gt = sl.shapes.add_table(len(rows), len(rows[0]), Inches(0.7), top, Inches(8.6), Inches(0.5) * len(rows))
            for ri, row in enumerate(rows):
                for ci, val in enumerate(row):
                    cell = gt.table.cell(ri, ci)
                    cell.text = val
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(18)
    prs.save(path)
    normalize_zip(path)


# -------------------------------------------------------------- XLSX / CSV
def tables_with_names(blocks):
    """(sheet name, rows) pairs: each table named after the preceding heading."""
    out, name = [], None
    for b in blocks:
        if b[0] == "h":
            name = b[2]
        elif b[0] == "table":
            out.append((name or f"Sheet{len(out) + 1}", b[1]))
    return out


def write_xlsx(blocks, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.properties.creator = wb.properties.lastModifiedBy = "Benchmark"
    wb.properties.created = wb.properties.modified = FIXED
    wb.remove(wb.active)
    for name, rows in tables_with_names(blocks):
        ws = wb.create_sheet(title=name[:31])
        for ri, row in enumerate(rows, start=1):
            for ci, val in enumerate(row, start=1):
                if val == "":
                    continue
                num = to_number(val) if ri > 1 else None
                cell = ws.cell(row=ri, column=ci, value=num if num is not None else val)
                if num is not None and "." in val:
                    cell.number_format = "0." + "0" * len(val.split(".")[1])
                if ri == 1:
                    cell.font = Font(bold=True)
    wb.save(path)
    normalize_zip(path)


def write_csv(rows, path):
    import csv
    with open(path, "w", encoding="utf-8", newline="") as f:
        csv.writer(f, lineterminator="\n").writerows(rows)


def table_md(name, rows):
    def line(r):
        return "| " + " | ".join(r) + " |"
    sep = "| " + " | ".join("---" for _ in rows[0]) + " |"
    return "\n".join([line(rows[0]), sep] + [line(r) for r in rows[1:]]) + "\n"


# -------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", help="comma-separated formats to (re)generate")
    ap.add_argument("--skip-pdf", action="store_true")
    args = ap.parse_args()
    only = set(args.only.split(",")) if args.only else None
    OUT.mkdir(exist_ok=True)
    derived = HERE / "ground_truth_derived"
    derived.mkdir(exist_ok=True)
    manifest = []

    def want(fmt):
        return only is None or fmt in only

    for stem, (lang, fmts) in PLAN.items():
        gt_file = f"ground_truth/{stem}.md"
        blocks = parse_md((GT / f"{stem}.md").read_text(encoding="utf-8"))
        for fmt in fmts:
            if fmt == "csv":
                for ti, (name, rows) in enumerate(tables_with_names(blocks), start=1):
                    fname = f"{stem}_t{ti}.csv"
                    dgt = derived / f"{stem}_t{ti}.md"
                    dgt.write_text(table_md(name, rows), encoding="utf-8")
                    if want("csv"):
                        write_csv(rows, OUT / fname)
                    manifest.append({"file": fname, "format": "csv", "gt": f"ground_truth_derived/{dgt.name}",
                                     "lang": lang, "carries": CARRIES["csv"]})
                continue
            fname = f"{stem}.{fmt}"
            if want(fmt) and not (fmt == "pdf" and args.skip_pdf):
                if fmt == "docx":
                    write_docx(blocks, OUT / fname, lang)
                elif fmt == "html":
                    (OUT / fname).write_text(render_html(blocks, lang, stem), encoding="utf-8")
                elif fmt == "pdf":
                    # render a throwaway copy of the HTML with headless Chrome
                    tmp_html = Path(tempfile.mkdtemp()) / f"{stem}.html"
                    tmp_html.write_text(render_html(blocks, lang, stem), encoding="utf-8")
                    write_pdf(tmp_html, (OUT / fname).resolve())
                    shutil.rmtree(tmp_html.parent)
                elif fmt == "pptx":
                    write_pptx(blocks, OUT / fname)
                elif fmt == "xlsx":
                    write_xlsx(blocks, OUT / fname)
            manifest.append({"file": fname, "format": fmt, "gt": gt_file, "lang": lang,
                             "carries": CARRIES[fmt]})
    manifest.sort(key=lambda m: (m["format"], m["file"]))
    (OUT / "manifest.json").write_text(json.dumps({"files": manifest}, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"wrote {len(manifest)} entries to {OUT / 'manifest.json'}")


if __name__ == "__main__":
    main()
